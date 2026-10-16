"""Helpers for the diagnostics template tests.

Reads the workbook's XML, works out what Excel shows for a cell, and fills the
template with generic values of the right type for each cell style.
"""

import datetime
import posixpath
import re
import zipfile
from xml.etree import ElementTree

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    diagnostics_cell_map as cell_map,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    diagnostics_workbook_utils as workbook_utils,
)

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

sheet_params = [pytest.param(s, id=s.name) for s in cell_map.DIAGNOSTICS_SHEETS]
charted_sheet_params = [
    pytest.param(s, id=s.name) for s in cell_map.DIAGNOSTICS_SHEETS if s.chart_count
]
comparison_sheet_params = [
    pytest.param(s, id=s.name)
    for s in (
        cell_map.COMP_SHORT_TERM,
        cell_map.COMP_MEDIUM_TERM,
        cell_map.COMP_LONG_TERM,
        cell_map.NO_FILTERS,
    )
]


# ---------------------------------------------------------------- cells
def cells_in(ranges: list[str]) -> set[tuple[int, int]]:
    """(row, column) of every cell in A1-style ranges."""
    return {
        (row, col)
        for cells in ranges
        for row, col in CellRange(cells).cells  # type: ignore[attr-defined]
    }


def coordinate(row: int, col: int) -> str:
    """A1 reference without creating the cell, which would hide its missing style."""
    return f"{get_column_letter(col)}{row}"


def effective_number_format(worksheet, row: int, col: int) -> str:
    """Number format Excel shows: the cell's own, else its column's.

    Use a freshly loaded workbook: reading a missing cell with `worksheet.cell()`
    creates it with the default style.
    """
    if workbook_utils.cell_exists(worksheet, row, col):
        return worksheet.cell(row, col).number_format
    dimension = workbook_utils.column_dimension(worksheet, col)
    return "General" if dimension is None else dimension.number_format


# ------------------------------------------------------- workbook XML
# An xlsx is a zip of XML parts. A part lists the parts it uses in
# `_rels/<part name>.rels`, as relationship id -> target path. To reach a chart,
# the tests walk workbook -> sheet -> drawing -> chart.
def read_xml(package: zipfile.ZipFile, part: str) -> ElementTree.Element:
    """Parses one part of the xlsx package."""
    return ElementTree.fromstring(package.read(part))


def _resolve(folder: str, target: str) -> str:
    """Package path of a relationship target.

    Example: ("xl/worksheets", "../drawings/drawing1.xml") gives
    "xl/drawings/drawing1.xml".
    """
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(folder, target))


def _relationships(package: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Relationship id -> package path of the part it points to, for one part."""
    folder, name = posixpath.split(part)
    rels_part = f"{folder}/_rels/{name}.rels"
    if rels_part not in package.namelist():
        return {}
    return {
        rel.get("Id"): _resolve(folder, rel.get("Target"))
        for rel in read_xml(package, rels_part).findall("rel:Relationship", NS)
    }


def _sheet_part(package: zipfile.ZipFile, sheet_name: str) -> str:
    """Package path of a sheet's XML part."""
    sheets = read_xml(package, "xl/workbook.xml").findall("main:sheets/main:sheet", NS)
    sheet_id = next(
        s.get(f"{{{NS['r']}}}id") for s in sheets if s.get("name") == sheet_name
    )
    return _relationships(package, "xl/workbook.xml")[sheet_id]


def _drawing_part(package: zipfile.ZipFile, sheet_part: str) -> str | None:
    """Package path of a sheet's drawing (the part that holds its charts), if any."""
    return next(
        (
            path
            for path in _relationships(package, sheet_part).values()
            if "/drawings/drawing" in path
        ),
        None,
    )


def drawing(
    package: zipfile.ZipFile, sheet_name: str
) -> tuple[ElementTree.Element, str]:
    """A sheet's drawing XML and its part name; an empty element if it has none."""
    part = _drawing_part(package, _sheet_part(package, sheet_name))
    if part is None:
        return ElementTree.Element("noDrawing"), ""
    return read_xml(package, part), part


def charts(package: zipfile.ZipFile, sheet_name: str) -> list[ElementTree.Element]:
    """The chart XML of every chart on a sheet."""
    sheet_drawing, drawing_part = drawing(package, sheet_name)
    chart_ids = [
        chart.get(f"{{{NS['r']}}}id")
        for chart in sheet_drawing.findall(".//c:chart", NS)
    ]
    if not chart_ids:
        return []
    chart_parts = _relationships(package, drawing_part)
    return [read_xml(package, chart_parts[rid]) for rid in chart_ids]


def split_reference(reference: str) -> tuple[str, str]:
    """('Sheet name', 'A1:B2') from a chart reference such as 'Sheet'!$A$1:$B$2."""
    sheet, cells = reference.rsplit("!", 1)
    return sheet.strip("'"), cells.replace("$", "")


def rows_of(reference: str) -> tuple[int, int]:
    """First and last row of a chart reference."""
    cells = CellRange(split_reference(reference)[1])
    return cells.min_row, cells.max_row


def row_count(reference: str) -> int:
    """Number of rows a chart reference covers."""
    first, last = rows_of(reference)
    return last - first + 1


def column_of(reference: str) -> str:
    """Column letters of a chart reference's first cell."""
    return re.match(r"[A-Z]+", split_reference(reference)[1]).group()


def axis_min(axis: ElementTree.Element) -> float | None:
    """A value axis's fixed minimum, or None when it is automatic."""
    found = axis.find("c:scaling/c:min", NS)
    return None if found is None else round(float(found.get("val")), 6)


def series_columns(chart: ElementTree.Element) -> set[str]:
    """Columns the chart's series values come from."""
    return {column_of(f.text) for f in chart.findall(".//c:ser/c:val//c:f", NS)}


def chart_title(chart: ElementTree.Element) -> str:
    """Chart title text; an empty string when there is no title."""
    title = chart.find("c:chart/c:title", NS)
    if title is None:
        return ""
    return "".join(run.text or "" for run in title.findall(".//a:t", NS))


def expected_comparison_title(columns: set[str]) -> str:
    """Title a comparison-sheet chart should have, from the columns it plots."""
    table = "care homes" if columns <= set("CDEF") else "non-residential"
    kind = "% change" if columns <= set("EFNO") else "Posts"
    return f"CQC {table} - {kind}"


def chart_signature(package: zipfile.ZipFile, sheet: cell_map.SheetMap) -> list:
    """What a sheet's charts plot and how they look.

    Series references, title, axis minimums, axis label positions and data label
    positions.
    """
    return [
        (
            [f.text for f in chart.findall(".//c:f", NS)],
            chart_title(chart),
            [axis_min(axis) for axis in chart.findall(".//c:valAx", NS)],
            [e.get("val") for e in chart.findall(".//c:tickLblPos", NS)],
            [e.get("val") for e in chart.findall(".//c:dLblPos", NS)],
        )
        for chart in charts(package, sheet.name)
    ]


# ------------------------------------------------------------- fixtures
FIRST_DATE = datetime.datetime(2026, 3, 1)


def add_months(start: datetime.datetime, months: int) -> datetime.datetime:
    """The date `months` months after `start`."""
    month_index = start.month - 1 + months
    return start.replace(
        year=start.year + month_index // 12, month=month_index % 12 + 1
    )


def generic_value(
    input_range: cell_map.InputRange, row_offset: int, col_offset: int
) -> object:
    """A value of the right type for a range's style, at an offset into the range.

    Text is the range's note, numbered down the rows, so each cell says what it
    is. Dates are consecutive months; counts and percentages rise down the rows
    and differ across the columns, so chart series do not sit on top of each other.
    """
    label = input_range.note or input_range.cells
    if CellRange(input_range.cells).size["rows"] > 1:
        label = f"{label} {row_offset + 1}"
    return {
        "text": label,
        "date": add_months(FIRST_DATE, row_offset),
        "count": 1000 * (row_offset + 1) + 100 * col_offset,
        "percent": round(0.01 * (row_offset + 1) + 0.002 * col_offset, 6),
    }[input_range.style]


def fill_inputs(
    workbook: openpyxl.Workbook, full: bool = False
) -> dict[tuple[str, str], object]:
    """Writes a generic value into the cell-map inputs; returns them by cell.

    `full` fills every slot. Otherwise only the first half of each range's rows
    are filled, so the copy also has empty slots.
    """
    written = {}
    for sheet in cell_map.DIAGNOSTICS_SHEETS:
        worksheet = workbook[sheet.name]
        for input_range in sheet.input_ranges:
            cells = CellRange(input_range.cells)
            rows = cells.size["rows"] if full else (cells.size["rows"] + 1) // 2
            for row in range(cells.min_row, cells.min_row + rows):
                for col in range(cells.min_col, cells.max_col + 1):
                    value = generic_value(
                        input_range, row - cells.min_row, col - cells.min_col
                    )
                    workbook_utils.write_input_cell(worksheet, row, col, value)
                    written[(sheet.name, coordinate(row, col))] = value
    return written
