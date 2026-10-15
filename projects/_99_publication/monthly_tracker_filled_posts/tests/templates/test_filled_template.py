"""A filled copy of the template saves with its values, charts and styles.

The copies use illustrative values. Set DIAGNOSTICS_TEMPLATE_CHECK_DIR to keep
them for a manual check in Excel: one filled to date, one with every slot filled,
and one for the first run after a year end.
"""

import os
import zipfile
from pathlib import Path

import openpyxl
import pytest

from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    diagnostics_cell_map as cell_map,
)
from projects._99_publication.monthly_tracker_filled_posts.tests.templates import (
    template_support as support,
)

CHECK_DIR_ENV = "DIAGNOSTICS_TEMPLATE_CHECK_DIR"


def _save_filled_copy(
    tmp_path_factory, file_name: str, full: bool, year_end: bool = False
):
    workbook = openpyxl.load_workbook(cell_map.TEMPLATE_PATH)
    written = support.fill_inputs(workbook, full=full, year_end=year_end)
    folder = Path(os.environ.get(CHECK_DIR_ENV) or tmp_path_factory.mktemp(file_name))
    path = folder / f"{file_name}.xlsx"
    workbook.save(path)
    return path, written


@pytest.fixture(scope="module")
def partly_filled(tmp_path_factory) -> tuple[Path, dict]:
    """Copy filled to date, as a real month would be."""
    return _save_filled_copy(tmp_path_factory, "diagnostics_filled_fixture", full=False)


@pytest.fixture(scope="module")
def fully_filled(tmp_path_factory) -> tuple[Path, dict]:
    """Copy with every slot filled, to check the formatting of all rows."""
    return _save_filled_copy(
        tmp_path_factory, "diagnostics_filled_fixture_full", full=True
    )


@pytest.fixture(scope="module")
def year_end_filled(tmp_path_factory) -> tuple[Path, dict]:
    """Copy for the first run after a year end, when the previous block is padded."""
    return _save_filled_copy(
        tmp_path_factory,
        "diagnostics_filled_fixture_year_end",
        full=False,
        year_end=True,
    )


def _comp_prev_dates(path: Path) -> tuple[list, list]:
    """Previous and current publication dates, row by row."""
    sheet = openpyxl.load_workbook(path)["Comp prev publication"]
    previous = [sheet.cell(row, 3).value for row in range(23, 60)]
    current = [sheet.cell(row, 3).value for row in range(64, 101)]
    return previous, current


def _misaligned_rows(path: Path) -> list[int]:
    """Rows where the previous date is filled but differs from the current date."""
    previous, current = _comp_prev_dates(path)
    return [
        row
        for row, (before, now) in enumerate(zip(previous, current))
        if before is not None and before != now
    ]


def _gap_rows(path: Path) -> list[int]:
    """Rows where the previous block is blank above a filled previous row."""
    previous, current = _comp_prev_dates(path)
    last_previous = max(
        (i for i, d in enumerate(previous) if d is not None), default=-1
    )
    return [
        row
        for row, (before, now) in enumerate(zip(previous, current))
        if before is None and now is not None and row < last_previous
    ]


class TestPartlyFilledCopy:
    def test_every_written_value_is_saved(self, partly_filled):
        """Values written to the inputs are the values read back."""
        path, written = partly_filled
        saved = openpyxl.load_workbook(path)

        unsaved = [
            key
            for key, value in written.items()
            if saved[key[0]][key[1]].value != value
        ]

        assert unsaved == [], "Written values that did not save"

    def test_comp_prev_previous_rows_line_up_with_current_rows(self, partly_filled):
        """The difference block compares row by row, so a previous row is blank or the same date."""
        assert _misaligned_rows(partly_filled[0]) == []

    def test_comp_prev_previous_block_has_no_gap_in_a_normal_month(self, partly_filled):
        """Only the latest month is missing from the previous block, at the end."""
        assert _gap_rows(partly_filled[0]) == []


class TestYearEndCopy:
    def test_comp_prev_previous_rows_line_up_with_current_rows(self, year_end_filled):
        """Padding keeps each previous row on the same row as the current row it compares to."""
        assert _misaligned_rows(year_end_filled[0]) == []

    def test_comp_prev_previous_block_is_padded_for_the_new_year(self, year_end_filled):
        """The current block has a row for the year just ended; the previous block has a blank one."""
        assert _gap_rows(year_end_filled[0]) != []


class TestFullyFilledCopy:
    def test_every_input_slot_is_filled(self, fully_filled):
        """The fixture reaches every input cell, so the style check covers them all."""
        written = fully_filled[1]
        empty = [
            f"{sheet.name}: {r.cells}: {support.coordinate(row, col)}"
            for sheet in cell_map.DIAGNOSTICS_SHEETS
            for r in sheet.input_ranges
            for row, col in sorted(support.cells_in([r.cells]))
            if (sheet.name, support.coordinate(row, col)) not in written
        ]

        assert empty == [], "Input cells the fixture did not fill"

    @pytest.mark.parametrize("sheet", support.sheet_params)
    def test_written_cells_keep_the_expected_style(
        self, sheet: cell_map.SheetMap, fully_filled
    ):
        """Written cells show the right font and number format, as the column's style."""
        worksheet = openpyxl.load_workbook(fully_filled[0])[sheet.name]
        wrong_style = [
            f"{r.cells}: {support.coordinate(row, col)}"
            for r in sheet.input_ranges
            for row, col in sorted(support.cells_in([r.cells]))
            if support.effective_style(worksheet, row, col)
            != cell_map.CELL_STYLES[r.style]
        ]

        assert wrong_style == [], "Written cells with the wrong font or number format"

    def test_charts_keep_what_they_plot_and_how_they_look(
        self, fully_filled, package: zipfile.ZipFile
    ):
        """Saving keeps each chart's series, title, axis minimums and label positions."""
        with zipfile.ZipFile(fully_filled[0]) as saved_package:
            changed = [
                sheet.name
                for sheet in cell_map.DIAGNOSTICS_SHEETS
                if support.chart_signature(saved_package, sheet)
                != support.chart_signature(package, sheet)
            ]

        assert changed == [], "Sheets whose charts changed on save"
