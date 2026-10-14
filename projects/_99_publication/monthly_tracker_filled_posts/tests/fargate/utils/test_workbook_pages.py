import re
import zipfile

import gptables as gpt
import openpyxl
import pytest

import projects._99_publication.monthly_tracker_filled_posts.fargate.utils.workbook_pages as job
from projects._99_publication.unittest_data.image_test_data import make_png
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils.workbook_text import (
    TABLE_TITLES,
)

PAGES = ["frontpage", "introduction", "revisions", "definitions", "contents"]
PLACEHOLDERS = {
    "latest_period": "August 2026",
    "publication_date": "September 15th 2026",
}


class TestLoadPageConfig:
    def test_config_has_a_section_for_each_of_the_five_pages(self):
        config = job.load_page_config()

        assert list(config) == PAGES
        assert [config[page]["sheet"] for page in PAGES] == [
            "Frontpage",
            "Introduction",
            "Revisions",
            "Definitions",
            "Contents",
        ]


class TestFillPlaceholders:
    @pytest.mark.parametrize("page", PAGES[:4])
    def test_config_text_placeholders_are_all_filled(self, page):
        rows = job.load_page_config()[page]["rows"]

        filled_rows = job.fill_placeholders(rows, PLACEHOLDERS)

        text = [
            value
            for row in filled_rows
            for value in row.values()
            if isinstance(value, str)
        ]
        assert not [
            value
            for value in text
            for name in PLACEHOLDERS
            if "{" + name + "}" in value
        ]

    def test_function_fills_known_placeholders(self):
        rows = [{"text": "Published {publication_date}"}, {"blank": 2}]

        assert job.fill_placeholders(rows, PLACEHOLDERS) == [
            {"text": "Published September 15th 2026"},
            {"blank": 2},
        ]

    def test_function_keeps_other_braces_as_text(self):
        rows = [{"text": "Use {braces} and {publication_date}"}]

        assert job.fill_placeholders(rows, PLACEHOLDERS) == [
            {"text": "Use {braces} and September 15th 2026"}
        ]

    def test_function_converts_non_string_text_to_a_string(self):
        rows = [{"text": 2026}, {"heading": True}]

        assert job.fill_placeholders(rows, PLACEHOLDERS) == [
            {"text": "2026"},
            {"heading": "True"},
        ]


class TestWritePages:
    @pytest.fixture(scope="class")
    def workbook(self, tmp_path_factory):
        path = tmp_path_factory.mktemp("pages") / "pages.xlsx"
        workbook = gpt.GPWorkbook(str(path))
        job.write_pages(workbook, PLACEHOLDERS, TABLE_TITLES)
        workbook.close()
        return openpyxl.load_workbook(path)

    @staticmethod
    def column_a(worksheet):
        return [cell.value for cell in worksheet["A"]]

    @staticmethod
    def blank_rows_before(worksheet, text):
        column_a = TestWritePages.column_a(worksheet)
        row = column_a.index(text)
        return [value is None for value in column_a[row - 2 : row]]

    def test_frontpage_shows_title_latest_period_and_publication_date(self, workbook):
        column_a = self.column_a(workbook["Frontpage"])

        assert column_a[0] == job.load_page_config()["frontpage"]["rows"][0]["heading"]
        assert "August 2026" in column_a
        assert "Publication Date: September 15th 2026" in column_a

    def test_frontpage_has_contact_details_and_address(self, workbook):
        column_a = self.column_a(workbook["Frontpage"])

        assert "Email: Analysis@skillsforcare.org.uk" in column_a
        assert "Telephone: 0113 245 1716" in column_a
        assert column_a[-5:] == [
            "Skills for Care",
            "West Gate",
            "6 Grace Street",
            "Leeds",
            "LS1 2RP",
        ]

    @pytest.mark.parametrize(
        "heading",
        [
            "Classification of Statistics",
            "Methodology and data coverage",
            "Uses of the report",
            "Useful links",
        ],
    )
    def test_introduction_has_two_blank_rows_before_each_heading(
        self, workbook, heading
    ):
        assert self.blank_rows_before(workbook["Introduction"], heading) == [True, True]

    def test_introduction_classification_paragraphs_are_separated_by_one_blank_row(
        self, workbook
    ):
        column_a = self.column_a(workbook["Introduction"])
        first = column_a.index("Classification of Statistics") + 1

        assert [value is None for value in column_a[first : first + 6]] == [
            False,
            True,
            False,
            True,
            False,
            True,
        ]

    def test_introduction_methodology_url_is_a_bare_link(self, workbook):
        cell = next(
            cell
            for cell in workbook["Introduction"]["A"]
            if str(cell.value).endswith("Filled-posts.aspx")
        )

        assert cell.hyperlink.target == cell.value

    def test_introduction_useful_links_are_descriptive_hyperlinks(self, workbook):
        introduction = workbook["Introduction"]
        first = self.column_a(introduction).index("Useful links") + 2
        cells = [introduction.cell(row, 1) for row in range(first, first + 6)]

        assert all(cell.hyperlink for cell in cells)
        assert not [cell.value for cell in cells if cell.value.startswith("http")]

    def test_revisions_states_original_publication_date(self, workbook):
        assert (
            "originally published on 15th July 2026"
            in workbook["Revisions"]["A2"].value
        )

    @pytest.mark.parametrize(
        "heading", ["Definitions", "Rounding", "Period", "Main service"]
    )
    def test_definitions_headings_are_bold(self, workbook, heading):
        definitions = workbook["Definitions"]
        row = self.column_a(definitions).index(heading) + 1

        assert definitions.cell(row, 1).font.b is True

    def test_definitions_rounding_url_is_a_bare_link(self, workbook):
        cell = next(
            cell
            for cell in workbook["Definitions"]["A"]
            if str(cell.value).endswith("Methodology.aspx")
        )

        assert cell.hyperlink.target == cell.value

    def test_definitions_text_is_not_wrapped(self, workbook):
        definitions = workbook["Definitions"]
        text_cells = [
            cell
            for cell in definitions["A"]
            if cell.value and not cell.font.b and not cell.hyperlink
        ]

        assert text_cells
        assert not [cell.value for cell in text_cells if cell.alignment.wrap_text]

    @pytest.mark.parametrize(
        "sheet,text_start",
        [
            ("Introduction", "This publication shows"),
            ("Frontpage", "Prepared by"),
            ("Revisions", "This document was originally"),
        ],
    )
    def test_other_pages_text_is_still_wrapped(self, workbook, sheet, text_start):
        cell = next(
            cell
            for cell in workbook[sheet]["A"]
            if str(cell.value).startswith(text_start)
        )

        assert cell.alignment.wrap_text is True

    def test_definitions_notation_key_defines_x_and_less_than_10(self, workbook):
        column_a = self.column_a(workbook["Definitions"])

        assert "[x] data is not published in this report." in column_a
        assert any(str(value).startswith("< 10 ") for value in column_a)

    def test_contents_lists_each_table_number_and_title_without_table_prefix(
        self, workbook
    ):
        contents = workbook["Contents"]

        assert [
            (row[0].value, row[1].value)
            for row in contents.iter_rows(min_row=4, max_row=6)
        ] == [(sheet, TABLE_TITLES[sheet]) for sheet in ["T0", "T1", "T2"]]
        assert [contents.cell(3, 1).value, contents.cell(3, 2).value] == [
            "Table number",
            "Table title",
        ]

    def test_contents_table_numbers_are_plain_text_not_links(self, workbook):
        cells = [workbook["Contents"].cell(row, 1) for row in range(4, 7)]

        assert [cell.value for cell in cells] == ["T0", "T1", "T2"]
        assert [cell.hyperlink for cell in cells] == [None, None, None]
        assert [cell.font.u for cell in cells] == [None, None, None]
        assert [cell.font.color.type for cell in cells] == ["theme"] * 3

    def test_contents_has_borders_on_every_cell(self, workbook):
        contents = workbook["Contents"]
        sides = [
            (
                cell.border.left.style,
                cell.border.right.style,
                cell.border.top.style,
                cell.border.bottom.style,
            )
            for row in contents.iter_rows(min_row=3, max_row=6, max_col=2)
            for cell in row
        ]

        assert set(sides) == {("thin", "thin", "thin", "thin")}

    @pytest.mark.parametrize(
        "sheet,expected_widths",
        [
            ("Frontpage", [76]),
            ("Introduction", [175]),
            ("Definitions", [120, 38]),
            ("Contents", [17, 86]),
        ],
    )
    def test_page_column_widths_match_reference(self, workbook, sheet, expected_widths):
        worksheet = workbook[sheet]
        widths = [
            worksheet.column_dimensions[letter].width
            for letter in "AB"[: len(expected_widths)]
        ]

        assert all(abs(w - e) < 1 for w, e in zip(widths, expected_widths))


class TestFrontpageLogo:
    @staticmethod
    def write(path, logo):
        workbook = gpt.GPWorkbook(str(path))
        job.write_pages(workbook, PLACEHOLDERS, TABLE_TITLES, logo=logo)
        workbook.close()
        return zipfile.ZipFile(path)

    @staticmethod
    def drawing(archive):
        return archive.read("xl/drawings/drawing1.xml").decode()

    def test_frontpage_logo_is_inserted_at_a5_when_supplied(self, tmp_path):
        archive = self.write(tmp_path / "logo.xlsx", make_png(200, 150))

        drawing = self.drawing(archive)
        frontpage_rels = archive.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        assert "drawing1.xml" in frontpage_rels
        assert re.search(
            r"<xdr:from><xdr:col>0</xdr:col><xdr:colOff>\d+</xdr:colOff>"
            r"<xdr:row>4</xdr:row>",
            drawing,
        )

    def test_frontpage_logo_is_marked_decorative(self, tmp_path):
        archive = self.write(tmp_path / "logo.xlsx", make_png(200, 150))

        assert re.search(r"<adec:decorative[^>]*val=\"1\"", self.drawing(archive))

    @pytest.mark.parametrize(
        "width,height,dpi",
        [
            (200, 150, 96),
            (50, 50, 96),
            (100, 100, 96),
            (400, 200, 96),
            (232, 347, 150),
        ],
        ids=["larger", "smaller", "exact", "wide", "high_dpi"],
    )
    def test_frontpage_logo_is_100_px_wide_at_any_size(
        self, tmp_path, width, height, dpi
    ):
        emu_per_px = 9525
        archive = self.write(tmp_path / "logo.xlsx", make_png(width, height, dpi))

        cx, cy = map(
            int,
            re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', self.drawing(archive)).groups(),
        )
        assert abs(cx - 100 * emu_per_px) <= emu_per_px
        assert abs(cy - 100 * height / width * emu_per_px) <= emu_per_px

    def test_no_logo_is_inserted_when_none_is_supplied(self, tmp_path):
        archive = self.write(tmp_path / "no_logo.xlsx", None)

        assert [
            name for name in archive.namelist() if "drawing" in name or "media" in name
        ] == []
