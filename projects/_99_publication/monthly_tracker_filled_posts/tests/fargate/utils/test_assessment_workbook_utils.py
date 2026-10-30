import datetime
import zipfile
from pathlib import Path

import boto3
import openpyxl
import polars as pl
import pytest
from moto import mock_aws
from openpyxl.styles import Font

from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    assessment_cell_map as cell_map,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    assessment_sheet_values as sheet_values,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    assessment_workbook_utils as job,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils.assessment_sheet_text import (
    TermFromDates,
    sheet_text,
)
from projects._99_publication.monthly_tracker_filled_posts.tests.templates import (
    template_support as support,
)
from utils.column_names.ind_cqc_pipeline_columns import IndCqcColumns as IndCQC
from utils.column_names.publication_columns import PublicationColumns as Pub
from utils.column_values.categorical_column_values import (
    PublishedJobGroupLabels,
    PublishedRegion,
)


APRIL_2024 = datetime.date(2024, 4, 1)
APRIL_2025 = datetime.date(2025, 4, 1)
APRIL_2026 = datetime.date(2026, 4, 1)
IMPORT_DATES = [APRIL_2024, APRIL_2025, APRIL_2026, datetime.date(2026, 5, 1)]
TERM_FROM_DATES = TermFromDates(short=APRIL_2026, medium=APRIL_2025, long=APRIL_2024)


def _assessment_lf(import_dates: list[datetime.date]) -> pl.LazyFrame:
    """One row per import date and service type, with every value column 1000."""
    value_columns = sheet_values.value_columns(cell_map.ASSESSMENT_SHEETS)
    rows = [
        {
            IndCQC.cqc_location_import_date: import_date,
            Pub.period_label_short: (import_date - datetime.timedelta(days=1)).strftime(
                "%b %Y"
            ),
            IndCQC.main_job_role_clean_labelled: PublishedJobGroupLabels.all_job_roles,
            IndCQC.current_region: PublishedRegion.england,
            IndCQC.primary_service_type: service_type,
        }
        | {column: 1000 for column in value_columns}
        for import_date in import_dates
        for service_type in sheet_values.SERVICE_TYPE_ORDER
    ]
    return pl.LazyFrame(rows)


def _sheet_with_arial_columns_b_to_d() -> openpyxl.worksheet.worksheet.Worksheet:
    worksheet = openpyxl.Workbook().active
    dimension = worksheet.column_dimensions["B"]
    dimension.min, dimension.max = 2, 4
    dimension.font = Font(name="Arial", sz=12)
    dimension.number_format = "#,##0"
    return worksheet


class TestCellExists:
    def test_is_false_until_the_cell_is_created(self):
        worksheet = openpyxl.Workbook().active

        before = job.cell_exists(worksheet, 2, 2)
        worksheet.cell(2, 2)

        assert (before, job.cell_exists(worksheet, 2, 2)) == (False, True)


class TestColumnDimension:
    def test_returns_the_dimension_covering_a_grouped_column(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        assert job.column_dimension(worksheet, 3) is worksheet.column_dimensions["B"]

    def test_returns_none_for_a_column_without_a_dimension(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        assert job.column_dimension(worksheet, 9) is None


class TestWriteInputCell:
    def test_gives_a_new_cell_its_column_style(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 3, "text")

        cell = worksheet.cell(3, 3)
        assert (cell.value, cell.font.name, cell.font.sz) == ("text", "Arial", 12)

    def test_gives_a_new_cell_its_column_number_format(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 3, 1234)

        assert worksheet.cell(3, 3).number_format == "#,##0"

    def test_leaves_a_new_cell_in_a_column_without_a_dimension_unstyled(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 9, "text")

        cell = worksheet.cell(3, 9)
        assert (cell.value, cell.font.name) == ("text", worksheet.cell(4, 9).font.name)

    def test_keeps_the_style_of_an_existing_cell(self):
        worksheet = _sheet_with_arial_columns_b_to_d()
        worksheet.cell(3, 3).font = Font(name="Calibri", sz=9)

        job.write_input_cell(worksheet, 3, 3, "text")

        assert worksheet.cell(3, 3).font.name == "Calibri"


PERIOD_LABELS = ["Mar 2026"]
RUN_NUMBER = 7


@pytest.fixture(scope="module")
def values() -> dict[str, dict[str, object]]:
    return support.generic_values(full=True)


@pytest.fixture(scope="module")
def filled_path(tmp_path_factory, values) -> Path:
    path = tmp_path_factory.mktemp("filled") / "assessment.xlsx"
    job.fill_assessment_workbook(
        PERIOD_LABELS, TERM_FROM_DATES, RUN_NUMBER, values, cell_map.TEMPLATE_PATH, path
    )
    return path


def _copy_template_with(tmp_path: Path, sheet: str, cell: str, value: object) -> Path:
    workbook = openpyxl.load_workbook(cell_map.TEMPLATE_PATH)
    workbook[sheet][cell] = value
    path = tmp_path / "changed_template.xlsx"
    workbook.save(path)
    return path


class TestFillAssessmentWorkbook:
    def test_values_land_in_the_mapped_cells(self, filled_path, values):
        saved = openpyxl.load_workbook(filled_path)

        wrong = [
            (sheet, cell)
            for sheet, cells in values.items()
            for cell, value in cells.items()
            if cell not in sheet_text(PERIOD_LABELS, TERM_FROM_DATES, RUN_NUMBER)[sheet]
            and saved[sheet][cell].value != value
        ]

        assert wrong == []

    def test_every_sheet_title_names_the_latest_period_and_run(self, tmp_path, values):
        output = tmp_path / "out.xlsx"

        job.fill_assessment_workbook(
            ["Sep 2025", "Mar 2026"],
            TERM_FROM_DATES,
            RUN_NUMBER,
            values,
            cell_map.TEMPLATE_PATH,
            output,
        )

        saved = openpyxl.load_workbook(output)
        titles = {
            sheet.name: saved[sheet.name][cell_map.TITLE_CELL].value
            for sheet in cell_map.ASSESSMENT_SHEETS
        }
        assert set(titles.values()) == {"End of March 2026 (run 7)"}

    def test_cells_outside_the_inputs_are_unchanged(self, filled_path):
        template = openpyxl.load_workbook(cell_map.TEMPLATE_PATH)
        saved = openpyxl.load_workbook(filled_path)

        changed = [
            (sheet.name, cell.coordinate)
            for sheet in cell_map.ASSESSMENT_SHEETS
            for row in template[sheet.name].iter_rows()
            for cell in row
            if (cell.row, cell.column)
            not in support.cells_in([r.cells for r in sheet.input_ranges])
            and saved[sheet.name][cell.coordinate].value != cell.value
        ]

        assert changed == []

    def test_saved_workbook_holds_no_formulas(self, filled_path):
        saved = openpyxl.load_workbook(filled_path)

        formulas = [
            (ws.title, cell.coordinate)
            for ws in saved
            for row in ws.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]

        assert formulas == []

    def test_charts_keep_count_and_series_references(self, filled_path):
        with (
            zipfile.ZipFile(cell_map.TEMPLATE_PATH) as before,
            zipfile.ZipFile(filled_path) as after,
        ):
            changed = [
                sheet.name
                for sheet in cell_map.ASSESSMENT_SHEETS
                if len(support.charts(after, sheet.name)) != sheet.chart_count
                or support.chart_signature(after, sheet)
                != support.chart_signature(before, sheet)
            ]

        assert changed == []

    def test_raises_when_a_value_targets_a_cell_outside_the_inputs(self, tmp_path):
        with pytest.raises(job.UnmappedCellError, match="'Filled posts'!B9"):
            job.fill_assessment_workbook(
                PERIOD_LABELS,
                TERM_FROM_DATES,
                RUN_NUMBER,
                {"Filled posts": {"B9": "overwritten"}},
                cell_map.TEMPLATE_PATH,
                tmp_path / "out.xlsx",
            )

    def test_raises_when_a_value_targets_an_unknown_sheet(self, tmp_path):
        with pytest.raises(job.UnmappedCellError, match="'Missing'!B3"):
            job.fill_assessment_workbook(
                PERIOD_LABELS,
                TERM_FROM_DATES,
                RUN_NUMBER,
                {"Missing": {"B3": "x"}},
                cell_map.TEMPLATE_PATH,
                tmp_path / "out.xlsx",
            )

    def test_raises_when_a_label_differs(self, tmp_path):
        template = _copy_template_with(tmp_path, "Filled posts", "C23", "Renamed")

        with pytest.raises(job.TemplateMismatchError, match="'Filled posts'!C23"):
            job.fill_assessment_workbook(
                PERIOD_LABELS,
                TERM_FROM_DATES,
                RUN_NUMBER,
                {},
                template,
                tmp_path / "out.xlsx",
            )

    def test_saves_nothing_when_error_raised(self, tmp_path):
        template = _copy_template_with(tmp_path, "Filled posts", "C23", "Renamed")
        output = tmp_path / "out.xlsx"

        with pytest.raises(job.TemplateMismatchError):
            job.fill_assessment_workbook(
                PERIOD_LABELS, TERM_FROM_DATES, RUN_NUMBER, {}, template, output
            )

        assert not output.exists()

    def test_raises_when_a_mapped_sheet_is_missing(self, tmp_path):
        workbook = openpyxl.load_workbook(cell_map.TEMPLATE_PATH)
        del workbook["No filters - long term"]
        template = tmp_path / "missing_sheet.xlsx"
        workbook.save(template)

        with pytest.raises(job.TemplateMismatchError, match="No filters - long term"):
            job.fill_assessment_workbook(
                PERIOD_LABELS,
                TERM_FROM_DATES,
                RUN_NUMBER,
                {},
                template,
                tmp_path / "out.xlsx",
            )


class TestFillAssessmentWorkbookFromLf:
    def test_comp_prev_labels_match_the_period_dates_beside_them(self, tmp_path):
        import_dates = [
            datetime.date(2024, 4, 1),
            datetime.date(2024, 5, 1),
            datetime.date(2025, 4, 1),
            datetime.date(2025, 10, 1),
            datetime.date(2026, 4, 1),
            datetime.date(2026, 5, 1),
            datetime.date(2026, 6, 1),
        ]
        output = tmp_path / "out.xlsx"

        job.fill_assessment_workbook_from_lf(
            _assessment_lf(import_dates),
            (cell_map.COMP_PREV_PUBLICATION,),
            TERM_FROM_DATES,
            RUN_NUMBER,
            cell_map.TEMPLATE_PATH,
            output,
        )

        sheet = openpyxl.load_workbook(output)[cell_map.COMP_PREV_PUBLICATION.name]
        for first_row in (23, 64, 106):
            rows = [
                (sheet[f"B{first_row + i}"].value, sheet[f"C{first_row + i}"].value)
                for i in range(6)
            ]
            assert rows == [
                ("Annual", datetime.datetime(2024, 3, 1)),
                ("Annual", datetime.datetime(2025, 3, 1)),
                ("Annual", datetime.datetime(2026, 3, 1)),
                ("Monthly", datetime.datetime(2026, 4, 1)),
                ("Monthly", datetime.datetime(2026, 5, 1)),
                (None, None),
            ]

    def test_saves_a_workbook_holding_the_filled_posts_values(self, tmp_path):
        output = tmp_path / "out.xlsx"

        job.fill_assessment_workbook_from_lf(
            _assessment_lf(IMPORT_DATES),
            (cell_map.FILLED_POSTS,),
            TERM_FROM_DATES,
            RUN_NUMBER,
            cell_map.TEMPLATE_PATH,
            output,
        )

        sheet = openpyxl.load_workbook(output)[cell_map.FILLED_POSTS.name]
        assert sheet["C24"].value == 1000

    def test_writes_the_run_number_into_every_sheet_title(self, tmp_path):
        output = tmp_path / "out.xlsx"

        job.fill_assessment_workbook_from_lf(
            _assessment_lf(IMPORT_DATES),
            (cell_map.FILLED_POSTS, cell_map.COMP_PREV_PUBLICATION),
            TERM_FROM_DATES,
            RUN_NUMBER,
            cell_map.TEMPLATE_PATH,
            output,
        )

        saved = openpyxl.load_workbook(output)
        titles = {
            sheet.name: saved[sheet.name][cell_map.TITLE_CELL].value
            for sheet in (cell_map.FILLED_POSTS, cell_map.COMP_PREV_PUBLICATION)
        }
        assert set(titles.values()) == {"End of April 2026 (run 7)"}

    def test_raises_for_a_sheet_without_a_values_builder(self, tmp_path):
        with pytest.raises(ValueError, match="No values builder"):
            job.fill_assessment_workbook_from_lf(
                _assessment_lf(IMPORT_DATES),
                (cell_map.SheetMap("Unknown sheet", (), 0),),
                TERM_FROM_DATES,
                RUN_NUMBER,
                cell_map.TEMPLATE_PATH,
                tmp_path / "out.xlsx",
            )


class TestUploadToS3:
    @mock_aws
    @pytest.mark.parametrize(
        "directory", ["s3://bucket/some/prefix/", "s3://bucket/some/prefix"]
    )
    def test_uploads_file_into_directory_with_its_name(self, tmp_path, directory):
        s3 = boto3.client("s3", region_name="eu-west-2")
        s3.create_bucket(
            Bucket="bucket",
            CreateBucketConfiguration={"LocationConstraint": "eu-west-2"},
        )
        local_file = tmp_path / "nested" / "summary.xlsx"
        local_file.parent.mkdir()
        local_file.write_bytes(b"workbook bytes")

        job.upload_to_s3(local_file, directory)

        body = s3.get_object(Bucket="bucket", Key="some/prefix/summary.xlsx")["Body"]
        assert body.read() == b"workbook bytes"
