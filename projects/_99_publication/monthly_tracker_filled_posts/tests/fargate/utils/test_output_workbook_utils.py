import zipfile
from datetime import date

import zipfile

import openpyxl
import polars as pl
import pytest

import projects._99_publication.monthly_tracker_filled_posts.fargate.utils.output_workbook_utils as job
import projects._99_publication.unittest_data.polars_pub_test_data as Data
from projects._99_publication.monthly_tracker_filled_posts.tests import xlsx_support
from projects._99_publication.unittest_data.image_test_data import make_png
from utils.column_names.publication_columns import PublicationColumns as Pub


class TestGetPublicationDate:
    @pytest.mark.parametrize(
        "case",
        [case.as_pytest_param() for case in Data.get_publication_date_test_cases],
    )
    def test_function_returns_expected_date(self, case):
        assert job.get_publication_date(case.latest_import_date) == case.expected


class TestFormatPublicationDate:
    @pytest.mark.parametrize(
        "case",
        [case.as_pytest_param() for case in Data.format_publication_date_test_cases],
    )
    def test_function_returns_expected_text(self, case):
        assert job.format_publication_date(case.publication_date) == case.expected


PUBLICATION_SCHEMA = pl.Schema(
    [
        (Pub.period, pl.Date),
        (Pub.period_label, pl.String),
        (Pub.region, pl.String),
        (Pub.main_service, pl.String),
        (Pub.estimated_filled_posts_rounded, pl.Float64),
        (Pub.cqc_locations, pl.Int64),
        (Pub.filled_posts_april_perc_change, pl.Float64),
        (Pub.filled_posts_cum_perc_change_short, pl.Float64),
        (Pub.locationid_count_april_perc_change, pl.Float64),
        (
            Pub.locationid_count_cum_perc_change_short,
            pl.Float64,
        ),
    ]
)


class TestPrepareT0ForPublication:
    input_schema = PUBLICATION_SCHEMA
    expected_columns = [
        Pub.period,
        Pub.region,
        Pub.main_service,
        Pub.estimated_filled_posts,
        Pub.cqc_locations,
    ]

    @pytest.mark.parametrize(
        "case",
        [case.as_pytest_param() for case in Data.prepare_t0_for_publication_test_cases],
    )
    def test_function_returns_expected_table(self, case):
        input_lf = pl.LazyFrame(case.input_data, self.input_schema, orient="row")

        returned_df = job.prepare_t0_for_publication(input_lf)

        assert list(returned_df.columns) == self.expected_columns
        assert returned_df.values.tolist() == [list(r) for r in case.expected_data]


class TestPreparePercChangeForPublication:
    input_schema = PUBLICATION_SCHEMA
    april_column = Pub.filled_posts_april_perc_change
    cumulative_column = Pub.filled_posts_cum_perc_change_short
    expected_columns = [
        Pub.period,
        Pub.region,
        Pub.main_service,
        Pub.annual_percentage_change,
        Pub.monthly_percentage_change,
    ]

    @pytest.mark.parametrize(
        "case",
        [
            case.as_pytest_param()
            for case in Data.prepare_perc_change_for_publication_test_cases
        ],
    )
    def test_function_returns_expected_table(self, case):
        input_lf = pl.LazyFrame(case.input_data, self.input_schema, orient="row")

        returned_df = job.prepare_perc_change_for_publication(
            input_lf, self.april_column, self.cumulative_column
        )

        assert list(returned_df.columns) == self.expected_columns
        assert returned_df.values.tolist() == [list(r) for r in case.expected_data]

    def test_function_returns_float64_without_float32_noise(self):
        input_lf = pl.LazyFrame(
            {
                Pub.period: [date(2026, 4, 1)],
                Pub.period_label: ["Mar 2026"],
                Pub.region: ["London"],
                Pub.main_service: ["All CQC locations"],
                self.april_column: pl.Series([0.1], dtype=pl.Float32),
                self.cumulative_column: pl.Series([0.0123456], dtype=pl.Float32),
            }
        )

        returned_df = job.prepare_perc_change_for_publication(
            input_lf, self.april_column, self.cumulative_column
        )

        assert returned_df[Pub.annual_percentage_change].tolist() == [0.1]
        assert returned_df[Pub.monthly_percentage_change].tolist() == [0.012]


class TestGetLatestPeriod:
    def test_function_returns_month_before_latest_period_date(self):
        publication_lf = pl.LazyFrame(
            Data.download_workbook_publication_data,
            PUBLICATION_SCHEMA,
            orient="row",
        )

        assert job.get_latest_period(publication_lf) == date(2026, 8, 1)


class TestBuildDownloadWorkbook:
    table_sheets = ["T0", "T1", "T2"]
    published_rows = len(Data.download_workbook_publication_data)

    @staticmethod
    def publication_lf():
        return pl.LazyFrame(
            Data.download_workbook_publication_data,
            PUBLICATION_SCHEMA,
            orient="row",
        )

    @pytest.fixture(scope="class")
    def workbook_path(self, tmp_path_factory):
        path = tmp_path_factory.mktemp("workbook") / "download.xlsx"

        job.build_download_workbook(self.publication_lf(), path)

        return path

    @pytest.fixture(scope="class")
    def workbook(self, workbook_path):
        return openpyxl.load_workbook(workbook_path)

    @staticmethod
    def column_a(worksheet):
        return [cell.value for cell in worksheet["A"]]

    @staticmethod
    def normalise(value):
        return " ".join(str(value).split())

    @staticmethod
    def header_row_number(worksheet):
        for cell in worksheet["A"]:
            if TestBuildDownloadWorkbook.normalise(cell.value).startswith(Pub.period):
                return cell.row

    def test_sheet_order_matches_the_reference(self, workbook):
        assert workbook.sheetnames == [
            "Frontpage",
            "Introduction",
            "Revisions",
            "Definitions",
            "Contents",
            "T0",
            "T1",
            "T2",
        ]

    @pytest.mark.parametrize(
        "sheet,expected_columns",
        [
            pytest.param(
                "T0",
                [
                    Pub.period,
                    Pub.region,
                    Pub.main_service,
                    Pub.estimated_filled_posts,
                    Pub.cqc_locations,
                ],
                id="T0",
            ),
            pytest.param(
                "T1",
                [
                    Pub.period,
                    Pub.region,
                    Pub.main_service,
                    Pub.annual_percentage_change,
                    Pub.monthly_percentage_change,
                ],
                id="T1",
            ),
            pytest.param(
                "T2",
                [
                    Pub.period,
                    Pub.region,
                    Pub.main_service,
                    Pub.annual_percentage_change,
                    Pub.monthly_percentage_change,
                ],
                id="T2",
            ),
        ],
    )
    def test_table_header_row_matches_columns(self, workbook, sheet, expected_columns):
        worksheet = workbook[sheet]
        header_row = worksheet[self.header_row_number(worksheet)]

        assert [
            self.normalise(cell.value) for cell in header_row if cell.value
        ] == expected_columns

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_table_row_count_matches_input(self, workbook, sheet):
        worksheet = workbook[sheet]
        first_data_row = self.header_row_number(worksheet) + 1

        data_rows = [
            cell for cell in worksheet["A"][first_data_row - 1 :] if cell.value
        ]

        assert len(data_rows) == self.published_rows

    @pytest.mark.parametrize(
        "sheet,title",
        [
            pytest.param("T0", job.T0_TITLE, id="T0"),
            pytest.param("T1", job.T1_TITLE, id="T1"),
            pytest.param("T2", job.T2_TITLE, id="T2"),
        ],
    )
    def test_table_has_title_subtitle_and_instruction_rows(
        self, workbook, sheet, title
    ):
        column_a = self.column_a(workbook[sheet])

        assert column_a[0] == title
        assert column_a[1] == "Skills for Care's workforce estimates as at Aug-2026"
        assert column_a[2] == "This sheet contains 1 table."

    def test_period_column_is_published_as_text_label(self, workbook):
        worksheet = workbook["T0"]

        assert (
            worksheet.cell(self.header_row_number(worksheet) + 1, 1).value == "Jul 2026"
        )

    @pytest.mark.parametrize(
        "sheet,expected_values",
        [("T1", (0.032, 0.005)), ("T2", (0.065, 0.011))],
    )
    def test_percentage_columns_use_0_0_percent_format(
        self, workbook, sheet, expected_values
    ):
        worksheet = workbook[sheet]
        # second data row has numeric annual and monthly changes
        row = self.header_row_number(worksheet) + 2
        annual, monthly = worksheet.cell(row, 4), worksheet.cell(row, 5)

        assert (annual.number_format, monthly.number_format) == ("0.0%", "0.0%")
        assert (annual.value, monthly.value) == expected_values

    def test_pages_and_tables_use_flat_12pt_arial(self, workbook):
        t0 = workbook["T0"]
        cells = [
            workbook["Frontpage"]["A1"],
            workbook["Introduction"]["A1"],
            workbook["Definitions"]["A1"],
            workbook["Contents"]["A3"],
            t0["A1"],
            t0["A2"],
            t0.cell(self.header_row_number(t0), 1),
        ]

        assert {(cell.font.name, cell.font.sz) for cell in cells} == {("Arial", 12.0)}

    def test_gridlines_are_hidden_on_every_sheet(self, workbook):
        assert [
            sheet.title
            for sheet in workbook.worksheets
            if sheet.sheet_view.showGridLines is not False
        ] == []

    def test_no_merged_cells_on_any_sheet(self, workbook):
        assert [sheet.title for sheet in workbook if sheet.merged_cells.ranges] == []

    def test_worksheet_title_is_in_cell_a1(self, workbook):
        assert all(sheet["A1"].value for sheet in workbook.worksheets)

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_table_header_is_preceded_by_a_blank_row(self, workbook, sheet):
        worksheet = workbook[sheet]
        above_header = worksheet[self.header_row_number(worksheet) - 1]

        assert [cell.value for cell in above_header] == [None] * len(above_header)

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_data_is_an_excel_table_over_header_and_data(self, workbook, sheet):
        worksheet = workbook[sheet]
        header_row = self.header_row_number(worksheet)

        assert [table.ref for table in worksheet.tables.values()] == [
            f"A{header_row}:E{header_row + self.published_rows}"
        ]

    def test_contents_table_has_no_filter_buttons(self, workbook):
        tables = workbook["Contents"].tables

        assert [(table.name, table.autoFilter) for table in tables.values()] == [
            ("contents_table", None)
        ]

    def test_table_sheets_hold_their_named_excel_tables(self, workbook):
        assert {name: list(workbook[name].tables) for name in self.table_sheets} == {
            "T0": ["t0_filled_posts"],
            "T1": ["t1_filled_posts_perc_change"],
            "T2": ["t2_location_count_perc_change"],
        }

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_no_sheet_level_filter_clashes_with_the_excel_table(self, workbook, sheet):
        assert workbook[sheet].auto_filter.ref is None

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_table_header_has_filter_buttons(self, workbook, sheet):
        tables = list(workbook[sheet].tables.values())

        assert [table.autoFilter is not None for table in tables] == [True]

    @pytest.mark.parametrize("sheet", table_sheets)
    def test_label_and_value_columns_are_boxed_separately(self, workbook, sheet):
        worksheet = workbook[sheet]
        first_data_row = self.header_row_number(worksheet) + 1
        borders = {
            column: (
                worksheet.cell(first_data_row, column).border.left.style,
                worksheet.cell(first_data_row, column).border.right.style,
            )
            for column in range(1, 6)
        }

        assert borders == {
            1: ("thin", None),
            2: (None, None),
            3: (None, "thin"),
            4: ("thin", None),
            5: (None, "thin"),
        }

    def test_value_columns_are_equally_wide_across_tables(self, workbook):
        widths = {
            name: (
                workbook[name].column_dimensions["D"].width,
                workbook[name].column_dimensions["E"].width,
            )
            for name in self.table_sheets
        }

        assert len(set(widths.values())) == 1

    def test_logo_is_added_to_the_frontpage(self, tmp_path):
        path = tmp_path / "with_logo.xlsx"

        job.build_download_workbook(
            self.publication_lf(), path, logo=make_png(200, 150)
        )

        archive = zipfile.ZipFile(path)
        assert "xl/media/image1.png" in archive.namelist()
        assert (
            "drawing1.xml"
            in archive.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        )

    def test_workbook_opens_on_the_frontpage_with_only_that_tab_selected(
        self, workbook
    ):
        assert workbook.active.title == "Frontpage"
        assert xlsx_support.selected_tab_titles(workbook) == ["Frontpage"]

    def test_every_sheet_has_cell_a1_selected(self, workbook):
        assert xlsx_support.sheets_not_at_a1(workbook) == []

    def test_build_download_workbook_runs_the_publication_data_once(self, tmp_path):
        runs = []

        def count(df):
            runs.append(1)
            return df

        publication_lf = self.publication_lf().map_batches(count)

        job.build_download_workbook(publication_lf, tmp_path / "once.xlsx")

        assert len(runs) == 1


class TestEnableTableFilters:
    table_xml = (
        '<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'id="1" name="{name}" displayName="{name}" ref="A5:E8" totalsRowShown="0">'
        "{filter}"
        '<tableColumns count="1"><tableColumn id="1" name="Period"/></tableColumns>'
        "</table>"
    )

    def write_zip(self, path, name="t0_filled_posts", table_filter=""):
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr(
                "xl/tables/table1.xml",
                self.table_xml.format(name=name, filter=table_filter),
            )
            archive.writestr("xl/workbook.xml", "<workbook/>")

    def test_function_adds_filter_buttons_to_the_named_tables(self, tmp_path):
        path = tmp_path / "workbook.xlsx"
        self.write_zip(path)

        job._enable_table_filters(path, ["t0_filled_posts"])

        xml = zipfile.ZipFile(path).read("xl/tables/table1.xml").decode()
        assert '<autoFilter ref="A5:E8"/><tableColumns' in xml

    def test_function_leaves_other_tables_without_filter_buttons(self, tmp_path):
        path = tmp_path / "workbook.xlsx"
        self.write_zip(path, name="contents_table")

        job._enable_table_filters(path, ["t0_filled_posts"])

        xml = zipfile.ZipFile(path).read("xl/tables/table1.xml").decode()
        assert "<autoFilter" not in xml

    def test_function_does_not_duplicate_an_existing_filter(self, tmp_path):
        path = tmp_path / "workbook.xlsx"
        self.write_zip(path, table_filter='<autoFilter ref="A5:E8"/>')

        job._enable_table_filters(path, ["t0_filled_posts"])

        xml = zipfile.ZipFile(path).read("xl/tables/table1.xml").decode()
        assert xml.count("<autoFilter") == 1

    def test_function_leaves_the_original_workbook_when_the_rewrite_fails(
        self, tmp_path, monkeypatch
    ):
        path = tmp_path / "workbook.xlsx"
        self.write_zip(path)
        original = path.read_bytes()

        def fail(*args, **kwargs):
            raise OSError("disk full")

        monkeypatch.setattr(zipfile.ZipFile, "writestr", fail)
        with pytest.raises(OSError):
            job._enable_table_filters(path, ["t0_filled_posts"])

        assert path.read_bytes() == original
        assert [file.name for file in tmp_path.iterdir()] == ["workbook.xlsx"]
