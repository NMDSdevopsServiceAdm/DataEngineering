import openpyxl
from openpyxl.styles import Font

from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    diagnostics_workbook_utils as job,
)


def _sheet_with_arial_columns_b_to_d() -> openpyxl.worksheet.worksheet.Worksheet:
    worksheet = openpyxl.Workbook().active
    dimension = worksheet.column_dimensions["B"]
    dimension.min, dimension.max = 2, 4
    dimension.font = Font(name="Arial", sz=12)
    dimension.number_format = "#,##0"
    return worksheet


class TestCellExists:
    def test_is_false_until_the_cell_is_created(self):
        worksheet = openpyxl.Workbook().active

        before = job.cell_exists(worksheet, 2, 2)
        worksheet.cell(2, 2)

        assert (before, job.cell_exists(worksheet, 2, 2)) == (False, True)


class TestColumnDimension:
    def test_returns_the_dimension_covering_a_grouped_column(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        assert job.column_dimension(worksheet, 3) is worksheet.column_dimensions["B"]

    def test_returns_none_for_a_column_without_a_dimension(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        assert job.column_dimension(worksheet, 9) is None


class TestWriteInputCell:
    def test_gives_a_new_cell_its_column_style(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 3, "text")

        cell = worksheet.cell(3, 3)
        assert (cell.value, cell.font.name, cell.font.sz) == ("text", "Arial", 12)

    def test_gives_a_new_cell_its_column_number_format(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 3, 1234)

        assert worksheet.cell(3, 3).number_format == "#,##0"

    def test_leaves_a_new_cell_in_a_column_without_a_dimension_unstyled(self):
        worksheet = _sheet_with_arial_columns_b_to_d()

        job.write_input_cell(worksheet, 3, 9, "text")

        cell = worksheet.cell(3, 9)
        assert (cell.value, cell.font.name) == ("text", worksheet.cell(4, 9).font.name)

    def test_keeps_the_style_of_an_existing_cell(self):
        worksheet = _sheet_with_arial_columns_b_to_d()
        worksheet.cell(3, 3).font = Font(name="Calibri", sz=9)

        job.write_input_cell(worksheet, 3, 3, "text")

        assert worksheet.cell(3, 3).font.name == "Calibri"
