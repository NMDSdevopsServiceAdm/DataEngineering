"""Writes values and text into the assessment template and saves it."""

from copy import copy
from pathlib import Path

import boto3
import openpyxl
import polars as pl
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.worksheet import Worksheet

from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    assessment_cell_map as cell_map,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils import (
    assessment_sheet_values as sheet_values,
)
from projects._99_publication.monthly_tracker_filled_posts.fargate.utils.assessment_sheet_text import (
    TermFromDates,
    sheet_text,
)
from utils.column_names.publication_columns import PublicationColumns as Pub
from utils.file_utils import split_s3_uri


class TemplateMismatchError(ValueError):
    """The template does not match the cell map."""


class UnmappedCellError(ValueError):
    """A value targets a sheet or cell outside the mapped input cells."""


def cell_exists(worksheet: Worksheet, row: int, col: int) -> bool:
    """Whether the sheet already holds a cell at this position.

    `worksheet.cell()` creates a missing cell, so existence is read from the
    private `_cells` mapping (openpyxl 3.1.5, pinned). A cell stored with the
    default style still counts as existing, which a style check cannot tell
    apart from a missing cell that falls back to its column's style.

    Args:
        worksheet (Worksheet): sheet to look in.
        row (int): 1-based row number.
        col (int): 1-based column number.

    Returns:
        bool: True if the cell exists.
    """
    return (row, col) in worksheet._cells


def column_dimension(worksheet: Worksheet, col: int) -> ColumnDimension | None:
    """Column dimension covering a column; openpyxl groups adjacent equal columns.

    Args:
        worksheet (Worksheet): sheet to look in.
        col (int): 1-based column number.

    Returns:
        ColumnDimension | None: the dimension, or None if the column has none.
    """
    return next(
        (d for d in worksheet.column_dimensions.values() if d.min <= col <= d.max),
        None,
    )


def write_input_cell(worksheet: Worksheet, row: int, col: int, value: object) -> None:
    """Writes a value, giving a new cell its column's style.

    Excel doesn't store a cell whose style equals its column's, and openpyxl
    doesn't apply a column's style to the cells it creates, so a new cell would
    otherwise get the workbook default font and number format.

    Args:
        worksheet (Worksheet): sheet to write to.
        row (int): 1-based row number.
        col (int): 1-based column number.
        value (object): value to write.
    """
    existed = cell_exists(worksheet, row, col)
    cell = worksheet.cell(row, col)
    dimension = column_dimension(worksheet, col)
    if not existed and dimension is not None:
        cell._style = copy(dimension._style)
    cell.value = value


def check_template(
    workbook: openpyxl.Workbook, values: dict[str, dict[str, object]]
) -> None:
    """Checks the template holds every mapped sheet and fixed label, and that
    every value targets a mapped input cell.

    Args:
        workbook (openpyxl.Workbook): the loaded template.
        values (dict[str, dict[str, object]]): values to write, keyed by sheet
            name, then by A1 cell reference.

    Raises:
        TemplateMismatchError: listing every missing sheet and label that differs.
        UnmappedCellError: if a value targets a sheet or cell outside the mapped
            input ranges.
    """
    problems = []
    for sheet in cell_map.ASSESSMENT_SHEETS:
        if sheet.name not in workbook.sheetnames:
            problems.append(f"missing sheet '{sheet.name}'")
            continue
        worksheet = workbook[sheet.name]
        for cell, expected in sheet.labels:
            found = worksheet[cell].value
            if found != expected:
                problems.append(
                    f"'{sheet.name}'!{cell} expected {expected!r}, found {found!r}"
                )
    if problems:
        raise TemplateMismatchError(
            "Template does not match cell map: " + "; ".join(problems)
        )

    input_cells = {
        sheet.name: _input_cells(sheet) for sheet in cell_map.ASSESSMENT_SHEETS
    }
    unmapped = [
        f"'{name}'!{cell}"
        for name, cells in values.items()
        for cell in cells
        if coordinate_to_tuple(cell) not in input_cells.get(name, set())
    ]
    if unmapped:
        raise UnmappedCellError(
            f"Values outside the mapped input cells: {', '.join(unmapped)}"
        )


def _input_cells(sheet: cell_map.SheetMap) -> set[tuple[int, int]]:
    return {
        (row, col)
        for input_range in sheet.input_ranges
        for row, col in CellRange(input_range.cells).cells
    }


def fill_assessment_workbook(
    period_labels: list[str],
    term_from_dates: TermFromDates,
    run_number: int,
    values: dict[str, dict[str, object]],
    template_path: Path,
    output_path: Path,
) -> None:
    """Fills the template with values and period text and saves a copy.

    Steps: build the period text with `sheet_text`, load the template, check it
    and the values with `check_template`, write the values and then the period
    text (which replaces any value for those cells), and save. Nothing is written
    if a check fails: `check_template` raises `TemplateMismatchError` or
    `UnmappedCellError`, and `sheet_text` raises `ValueError` for empty period
    labels or none in March.

    Args:
        period_labels (list[str]): the run's period labels, such as "Mar 2026".
        term_from_dates (TermFromDates): first import date of each comparison term.
        run_number (int): the run the workbook is built from, shown in the titles.
        values (dict[str, dict[str, object]]): values to write, keyed by sheet
            name, then by A1 cell reference, e.g. `{"Filled posts": {"C24":
            1000}}`. Cells must be in the sheet's mapped input ranges.
        template_path (Path): template to load.
        output_path (Path): local path to save to.
    """
    text = sheet_text(period_labels, term_from_dates, run_number)
    workbook = openpyxl.load_workbook(template_path)
    check_template(workbook, values)

    for name, cells in values.items():
        for cell, value in cells.items():
            row, col = coordinate_to_tuple(cell)
            write_input_cell(workbook[name], row, col, value)
    for name, cells in text.items():
        for cell, value in cells.items():
            row, col = coordinate_to_tuple(cell)
            write_input_cell(workbook[name], row, col, value)
    workbook.save(output_path)


def fill_assessment_workbook_from_lf(
    assessment_lf: pl.LazyFrame,
    sheets: tuple[cell_map.SheetMap, ...],
    term_from_dates: TermFromDates,
    run_number: int,
    template_path: Path,
    output_path: Path,
) -> None:
    """Writes the assessment data into the given sheets of the template.

    The workbook is written once, so every sheet's values are built first.

    Args:
        assessment_lf (pl.LazyFrame): assessment data.
        sheets (tuple[cell_map.SheetMap, ...]): sheets to fill.
        term_from_dates (TermFromDates): first import date of each comparison term.
        run_number (int): the run the data was built from, shown in the titles.
        template_path (Path): template to load.
        output_path (Path): local path to save to.

    Raises:
        ValueError: if a sheet has no values builder yet.
    """
    builders = sheet_values.SHEET_VALUE_BUILDERS
    unsupported = [s.name for s in sheets if s.name not in builders]
    if unsupported:
        raise ValueError(f"No values builder for sheets: {', '.join(unsupported)}")

    assessment_df = sheet_values.prepare_assessment_df(
        assessment_lf, sheet_values.value_columns(sheets)
    )
    values = {
        sheet.name: builders[sheet.name](assessment_df, term_from_dates)
        for sheet in sheets
    }
    period_labels = assessment_df.get_column(Pub.period_label_short).unique().to_list()
    fill_assessment_workbook(
        period_labels, term_from_dates, run_number, values, template_path, output_path
    )


def upload_to_s3(local_path: Path, s3_directory: str) -> None:
    """Uploads a local file into an S3 directory, keeping its file name.

    Args:
        local_path (Path): file to upload.
        s3_directory (str): S3 directory to put the file in.
    """
    bucket, prefix = split_s3_uri(s3_directory)
    key = f"{prefix.rstrip('/')}/{local_path.name}"
    boto3.client("s3").upload_file(str(local_path), bucket, key)
